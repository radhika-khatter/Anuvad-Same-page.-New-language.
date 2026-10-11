"""
Scanned English PDF -> Hindi Word (.docx)

Pipeline: PDF page -> image -> OCR (Tesseract) -> clean paragraphs
          -> translate -> Word file

Engines (--engine):
    google : free Google endpoint via deep-translator (can get rate-limited)
    nllb   : Meta NLLB model, runs on your PC, no limits (needs: pip install torch transformers sentencepiece)
    argos  : Argos Translate, runs on your PC, lighter but more literal (needs: pip install argostranslate)

Usage:
    python pdf_translate_hi.py input.pdf --engine nllb
    python pdf_translate_hi.py input.pdf --engine nllb --start 1 --end 3     # test a few pages

Windows (if tools are not on PATH):
    python pdf_translate_hi.py input.pdf --engine nllb ^
        --tesseract "C:\\Program Files\\Tesseract-OCR\\tesseract.exe" ^
        --poppler "C:\\poppler\\Library\\bin"

Progress is saved to <name>_progress.json. Re-run the same command to resume.
"""
import argparse
import json
import re
import sys
import time
from pathlib import Path

import pytesseract
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pdf2image import convert_from_path, pdfinfo_from_path
from PIL import ImageFilter, ImageOps

HINDI_FONT = "Mangal"  # ships with Windows; use "Noto Sans Devanagari" if installed
MAX_CHARS = 4500       # Google free endpoint limit is 5000
PAUSE = 1.5            # seconds between Google requests


# ---------- OCR ----------
def preprocess(img):
    img = ImageOps.grayscale(img)
    img = ImageOps.autocontrast(img)
    img = img.filter(ImageFilter.SHARPEN)
    return img


def ocr_page(img, lang="eng"):
    return pytesseract.image_to_string(preprocess(img), lang=lang, config="--psm 3")


def clean_to_paragraphs(text):
    """Turn raw OCR text into a list of paragraphs."""
    text = text.replace("\r", "")
    text = re.sub(r"-\n(?=[a-z])", "", text)  # join hyphenated line breaks
    paras = []
    for block in re.split(r"\n\s*\n", text):
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        # drop lone page numbers
        lines = [ln for ln in lines if not re.fullmatch(r"[-\u2013 ]*\d{1,4}[-\u2013 ]*", ln)]
        if not lines:
            continue
        para = re.sub(r"\s+", " ", " ".join(lines)).strip()
        if para:
            paras.append(para)
    return paras


# ---------- Translation engines ----------
def split_long(text, limit=MAX_CHARS):
    if len(text) <= limit:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    chunks, cur = [], ""
    for s in sentences:
        while len(s) > limit:  # absurdly long sentence
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(s[:limit])
            s = s[limit:]
        if len(cur) + len(s) + 1 > limit:
            chunks.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        chunks.append(cur)
    return chunks


def sentence_groups(text, limit=300):
    """Small groups of sentences: local models translate these best."""
    sents = re.split(r"(?<=[.!?])\s+", text)
    groups, cur = [], ""
    for s in sents:
        if cur and len(cur) + len(s) + 1 > limit:
            groups.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        groups.append(cur)
    return groups


class GoogleEngine:
    WAITS = [10, 30, 60, 120, 240]

    def __init__(self):
        from deep_translator import GoogleTranslator
        self.t = GoogleTranslator(source="en", target="hi")

    def translate(self, text):
        out = []
        for chunk in split_long(text):
            for wait in self.WAITS + [None]:
                try:
                    out.append(self.t.translate(chunk) or "")
                    break
                except Exception as e:
                    if wait is None:
                        raise RuntimeError(
                            "Google keeps refusing requests (rate limit / block). "
                            "Wait a while, or re-run with --engine nllb."
                        ) from e
                    print(f"    Google refused ({type(e).__name__}); waiting {wait}s...")
                    time.sleep(wait)
            time.sleep(PAUSE)
        return " ".join(out)


class NLLBEngine:
    MODEL = "facebook/nllb-200-distilled-600M"

    def __init__(self):
        try:
            import torch
            from transformers import AutoModelForSeq2SeqLM, AutoTokenizer
        except ImportError:
            sys.exit("Run:  pip install torch transformers sentencepiece")
        print("Loading NLLB model (first run downloads ~2.5 GB, please wait)...")
        self.torch = torch
        self.tok = AutoTokenizer.from_pretrained(self.MODEL, src_lang="eng_Latn")
        self.model = AutoModelForSeq2SeqLM.from_pretrained(self.MODEL)
        self.model.eval()
        if torch.cuda.is_available():
            self.device = "cuda"
            self.model.to("cuda").half()  # fp16: faster, ~1.5 GB of GPU memory
            self.batch = 8
            print(f"NLLB running on GPU: {torch.cuda.get_device_name(0)}")
        else:
            self.device = "cpu"
            self.batch = 4
            print("NLLB running on CPU (CUDA PyTorch not found, so no GPU is used)")
        self.hin_id = self.tok.convert_tokens_to_ids("hin_Deva")

    def translate(self, text):
        groups = sentence_groups(text)
        outs = []
        for i in range(0, len(groups), self.batch):  # small batches
            batch = groups[i:i + self.batch]
            enc = self.tok(batch, return_tensors="pt", padding=True, truncation=True, max_length=512)
            enc = enc.to(self.device)
            with self.torch.no_grad():
                gen = self.model.generate(
                    **enc, forced_bos_token_id=self.hin_id, num_beams=4, max_new_tokens=600
                )
            outs.extend(self.tok.batch_decode(gen, skip_special_tokens=True))
        return " ".join(outs)


class ArgosEngine:
    def __init__(self):
        try:
            import argostranslate.package as pkg
            import argostranslate.translate as tr
        except ImportError:
            sys.exit("Run:  pip install argostranslate")
        self.tr = tr
        langs = {l.code: l for l in tr.get_installed_languages()}
        if not ("en" in langs and "hi" in langs and langs["en"].get_translation(langs["hi"])):
            print("Downloading Argos English->Hindi model (one time)...")
            pkg.update_package_index()
            found = next(
                (p for p in pkg.get_available_packages() if p.from_code == "en" and p.to_code == "hi"),
                None,
            )
            if not found:
                sys.exit("Argos has no English->Hindi package right now. Use --engine nllb.")
            pkg.install_from_path(found.download())

    def translate(self, text):
        return " ".join(self.tr.translate(g, "en", "hi") for g in sentence_groups(text))


def make_engine(name):
    return {"google": GoogleEngine, "nllb": NLLBEngine, "argos": ArgosEngine}[name]()


# ---------- Word output ----------
def set_complex_font(style, font_name, size_pt):
    """Apply a Devanagari-capable font (Word needs the 'cs' font slot)."""
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size_pt * 2)))


def build_docx(pages, out_path):
    doc = Document()
    set_complex_font(doc.styles["Normal"], HINDI_FONT, 12)
    for i, (page_no, paras) in enumerate(pages):
        if i > 0:
            doc.add_page_break()
        h = doc.add_paragraph()
        r = h.add_run(f"\u092a\u0943\u0937\u094d\u0920 {page_no}")  # "पृष्ठ N"
        r.bold = True
        for p in paras:
            doc.add_paragraph(p)
    doc.save(out_path)


# ---------- Main ----------
def main():
    ap = argparse.ArgumentParser(description="Translate a scanned English PDF into a Hindi Word file")
    ap.add_argument("pdf")
    ap.add_argument("-o", "--output")
    ap.add_argument("--engine", choices=["google", "nllb", "argos"], default="google")
    ap.add_argument("--dpi", type=int, default=300)
    ap.add_argument("--start", type=int, default=1)
    ap.add_argument("--end", type=int, default=None)
    ap.add_argument("--lang", default="eng", help="Tesseract language (default eng)")
    ap.add_argument("--tesseract", help="Path to tesseract.exe (Windows)")
    ap.add_argument("--poppler", help="Path to poppler 'bin' folder (Windows)")
    args = ap.parse_args()

    pdf = Path(args.pdf)
    if not pdf.exists():
        sys.exit(f"File not found: {pdf}")
    out_path = Path(args.output) if args.output else pdf.with_name(pdf.stem + "_hindi.docx")
    progress_path = pdf.with_name(pdf.stem + "_progress.json")
    english_path = pdf.with_name(pdf.stem + "_english_ocr.txt")

    if args.tesseract:
        pytesseract.pytesseract.tesseract_cmd = args.tesseract

    poppler = args.poppler
    total = pdfinfo_from_path(str(pdf), poppler_path=poppler)["Pages"]
    end = min(args.end or total, total)
    print(f"PDF has {total} pages. Processing {args.start}-{end}.")

    progress = {}
    if progress_path.exists():
        progress = json.loads(progress_path.read_text(encoding="utf-8"))
        print(f"Resuming: {len(progress)} pages already done.")

    engine = make_engine(args.engine)

    for n in range(args.start, end + 1):
        if str(n) in progress:
            continue
        print(f"Page {n}/{end}: OCR...")
        img = convert_from_path(
            str(pdf), dpi=args.dpi, first_page=n, last_page=n, poppler_path=poppler
        )[0]
        paras = clean_to_paragraphs(ocr_page(img, args.lang))
        print(f"Page {n}/{end}: translating {len(paras)} paragraphs...")
        try:
            hindi = [engine.translate(p) for p in paras]
        except RuntimeError as e:
            print(f"\nStopped on page {n}: {e}")
            print("Pages finished so far are saved. Re-run the same command to resume.")
            break
        progress[str(n)] = {"en": paras, "hi": hindi}
        progress_path.write_text(json.dumps(progress, ensure_ascii=False, indent=1), encoding="utf-8")

    # Outputs (built from whatever pages are finished)
    page_nums = sorted(int(k) for k in progress if args.start <= int(k) <= end)
    if not page_nums:
        sys.exit("No pages were translated, nothing to write.")
    build_docx([(n, progress[str(n)]["hi"]) for n in page_nums], out_path)
    with open(english_path, "w", encoding="utf-8") as f:
        for n in page_nums:
            f.write(f"\n=== Page {n} ===\n")
            f.write("\n\n".join(progress[str(n)]["en"]) + "\n")

    print(f"\nDone ({len(page_nums)} pages).\n  Hindi Word file : {out_path}\n  English OCR text: {english_path}")
    print("Tip: compare the OCR text with your PDF to spot scan errors.")


if __name__ == "__main__":
    main()